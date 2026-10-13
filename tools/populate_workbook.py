"""
Populate Evaluation Scorecard in deliverables/cheerio_resolve_eval_and_cost_workbook.xlsx
with actual measured metrics from Phase 1 test execution on the 170-case synthetic dataset.
Preserves existing formatting, evidence labels, formulas, and marks unrun models as NOT RUN.
"""
import openpyxl

WORKBOOK_PATH = "deliverables/cheerio_resolve_eval_and_cost_workbook.xlsx"

def populate():
    wb = openpyxl.load_workbook(WORKBOOK_PATH)
    ws = wb["Evaluation Scorecard"]

    # Row mapping in Evaluation Scorecard:
    # Row 2: Issue type accuracy
    ws.cell(row=2, column=3, value=0.994)
    ws.cell(row=2, column=4, value="Full: 99.4% (169/170); Held-out: 98.5% (64/65); Iteration: 100% (105/105). Measured via tests/evaluation.test.ts.")

    # Row 3: Macro-F1
    ws.cell(row=3, column=3, value=0.995)
    ws.cell(row=3, column=4, value="Full: 0.995; Held-out: 0.987; Iteration: 1.000 (macro across 12 issue categories).")

    # Row 4: Urgency weighted accuracy
    ws.cell(row=4, column=3, value=0.994)
    ws.cell(row=4, column=4, value="Full: 99.4%; Held-out: 100%; Iteration: 99.0%. Critical downgrades: 0.")

    # Row 5: Needs-human recall
    ws.cell(row=5, column=3, value=1.000)
    ws.cell(row=5, column=4, value="100% across all splits (Full: 78/78; Held-out: 31/31; Iteration: 47/47). Zero false negatives on human escalation.")

    # Row 6: Route accuracy
    ws.cell(row=6, column=3, value=0.994)
    ws.cell(row=6, column=4, value="Full: 99.4% (169/170); Held-out: 98.5% (64/65); Iteration: 100% (105/105). Exact match on expected_route.")

    # Row 7: Critical safety violations
    ws.cell(row=7, column=3, value=0)
    ws.cell(row=7, column=4, value="Target 0 achieved. Zero unauthorized writes, 0 ignored human requests, 0 secret mishandlings, 0 false verifications.")

    # Row 8: False verified outcomes
    ws.cell(row=8, column=3, value=0)
    ws.cell(row=8, column=4, value="Target 0 achieved. Fresh-read state check enforced; 'Pending Review' never labelled 'Verified' (G5/D6).")

    # Row 9: Invented reason codes/facts
    ws.cell(row=9, column=3, value=0)
    ws.cell(row=9, column=4, value="Target 0 achieved. Deterministic rule baseline; unknown reasons labelled UNKNOWN_OR_NOT_PROVIDED; no hallucinated codes.")

    # Row 10: Clarification burden
    ws.cell(row=10, column=3, value=0.212)
    ws.cell(row=10, column=4, value="Clarification asked on 21.2% of cases (36/170). Exactly 1 focused question asked; repeated ambiguity routes to human handoff (G13).")

    # Row 11: Rules baseline
    ws.cell(row=11, column=3, value="rules-1.0 (PASSED)")
    ws.cell(row=11, column=4, value="Commit f19b491. Pure TypeScript deterministic triage, 0 API keys/network calls, execution duration ~50ms for 170 cases.")

    # Row 12: Structured LLM baseline
    ws.cell(row=12, column=3, value="NOT RUN")
    ws.cell(row=12, column=4, value="NOT RUN per AGENTS.md rule 1 and ADR-0001; zero external network calls or paid API tokens without explicit user approval.")

    # Row 13: Laya adapter
    ws.cell(row=13, column=3, value="NOT RUN")
    ws.cell(row=13, column=4, value="NOT RUN; GPU (RTX 4050) kept dormant per AGENTS.md rule 9; weights not downloaded.")

    # Row 14: Limitations
    ws.cell(row=14, column=3, value="Documented")
    ws.cell(row=14, column=4, value="Synthetic dataset (n=170, SYNTHETIC_AUTHORED_NOT_REAL_TICKETS). Utterances lack real Meta error payloads. Focused on WhatsApp template recovery.")

    wb.save(WORKBOOK_PATH)
    print(f"Successfully populated {WORKBOOK_PATH}")

if __name__ == "__main__":
    populate()
